from __future__ import annotations

import logging
from pathlib import Path

from app.models.entities import LoadedDocument
from app.services.text_cleaner import TextCleaner

logger = logging.getLogger(__name__)


class DocumentLoader:
    def __init__(self, allowed_extensions: tuple[str, ...], cleaner: TextCleaner | None = None) -> None:
        self.allowed_extensions = {ext.lower() for ext in allowed_extensions}
        self.cleaner = cleaner or TextCleaner()

    def discover_files(self, source_dir: Path, recursive: bool = True) -> list[Path]:
        if not source_dir.exists():
            raise FileNotFoundError(f"Source directory does not exist: {source_dir}")
        if not source_dir.is_dir():
            raise NotADirectoryError(f"Expected a directory path: {source_dir}")

        iterator = source_dir.rglob("*") if recursive else source_dir.glob("*")
        files = [
            path
            for path in iterator
            if path.is_file() and path.suffix.lower() in self.allowed_extensions
        ]
        return sorted(files)

    def load_file(self, file_path: Path) -> LoadedDocument:
        extension = file_path.suffix.lower()
        metadata: dict[str, object] = {
            "extension": extension,
            "filename": file_path.name,
            "size_bytes": file_path.stat().st_size,
            "source_filename": file_path.name,
        }

        if extension in {".txt", ".md"}:
            raw_text = file_path.read_text(encoding="utf-8", errors="ignore")
            text = self.cleaner.clean(raw_text)
        elif extension == ".pdf":
            pages = self._read_pdf_pages(file_path)
            text, page_map = self._merge_pages(pages)
            metadata["page_count"] = len(page_map)
            metadata["page_map"] = page_map
        elif extension == ".docx":
            raw_text = self._read_docx(file_path)
            text = self.cleaner.clean(raw_text)
        else:
            raise ValueError(f"Unsupported file type: {extension}")

        if not text:
            raise ValueError(f"No readable text found in {file_path.name}")

        return LoadedDocument(
            source=str(file_path),
            text=text,
            metadata=metadata,
        )

    def _read_pdf_pages(self, file_path: Path) -> list[str]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ImportError(
                "PDF support requires 'pypdf'. Install dependencies from requirements.txt."
            ) from exc

        reader = PdfReader(str(file_path))
        pages: list[str] = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            cleaned_page = self.cleaner.clean(page_text)
            pages.append(cleaned_page)
        return pages

    def _read_docx(self, file_path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError(
                "DOCX support requires 'python-docx'. Install dependencies from requirements.txt."
            ) from exc

        doc = Document(str(file_path))
        paragraphs = [paragraph.text for paragraph in doc.paragraphs]
        return "\n".join(paragraphs)

    def _merge_pages(self, pages: list[str]) -> tuple[str, list[dict[str, int]]]:
        merged_parts: list[str] = []
        page_map: list[dict[str, int]] = []
        cursor = 0

        for page_index, page_text in enumerate(pages, start=1):
            normalized_page = page_text.strip()
            if not normalized_page:
                continue

            if merged_parts:
                merged_parts.append("\n\n")
                cursor += 2

            start_char = cursor
            merged_parts.append(normalized_page)
            cursor += len(normalized_page)
            end_char = cursor
            page_map.append(
                {
                    "page_number": page_index,
                    "start_char": start_char,
                    "end_char": end_char,
                }
            )

        return "".join(merged_parts).strip(), page_map

